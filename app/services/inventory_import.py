from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from typing import Any

from app.providers.taobao_provider import parse_taobao_products
from scoring import InventoryItem, parse_inventory


@dataclass(frozen=True)
class SmartInventoryRow:
    product_name: str
    sku: str = ""
    color: str = ""
    cost_price: float | None = None
    cost_currency: str = "CAD"
    inventory: int | None = None
    target_price: float | None = None
    notes: str = ""


def build_inventory_items(
    inventory_text: str,
    taobao_json_text: str = "",
    excel_bytes: bytes | None = None,
    excel_filename: str = "",
) -> list[InventoryItem]:
    items: list[InventoryItem] = []
    if excel_bytes:
        items.extend(smart_rows_to_inventory_items(parse_smart_inventory_file(excel_bytes, excel_filename)))
    if inventory_text.strip():
        items.extend(parse_inventory(inventory_text))
    if items:
        return _dedupe_inventory_items(items)

    if taobao_json_text.strip():
        taobao_items = parse_taobao_inventory_json(taobao_json_text)
        if taobao_items:
            return _dedupe_inventory_items(taobao_items)
        raise ValueError("淘宝 JSON 中没有 targetProductStatus = 1 的有效商品。")

    if not items:
        raise ValueError("请至少上传 Excel、粘贴淘宝 JSON，或填写一个库存商品。")
    return _dedupe_inventory_items(items)


def parse_excel_inventory(file_bytes: bytes, filename: str = "") -> list[InventoryItem]:
    return smart_rows_to_inventory_items(parse_smart_inventory_file(file_bytes, filename))


def parse_smart_inventory_file(file_bytes: bytes, filename: str = "") -> list[SmartInventoryRow]:
    rows = _load_table_rows(file_bytes, filename)
    if not rows:
        return []
    header_index = _detect_header_row(rows)
    if header_index is None:
        return _parse_headerless_inventory_rows(rows)

    headers = [_cell_text(value) for value in rows[header_index]]
    mapping = _map_headers(headers, rows[header_index + 1 : header_index + 8])
    if "product_name" not in mapping:
        raise ValueError("没有识别到商品名列。支持表头：商品名、品名、product、title、name。")

    parsed: list[SmartInventoryRow] = []
    for raw_row in rows[header_index + 1 :]:
        if _is_empty_row(raw_row):
            continue
        product_name = _cell_from_mapping(raw_row, mapping, "product_name")
        if not product_name:
            continue
        cost_text = _cell_from_mapping(raw_row, mapping, "cost_price")
        parsed.append(
            SmartInventoryRow(
                product_name=product_name,
                sku=_cell_from_mapping(raw_row, mapping, "sku"),
                color=_cell_from_mapping(raw_row, mapping, "color"),
                cost_price=_optional_money_text(cost_text),
                cost_currency=_infer_cost_currency(cost_text, _cell_from_mapping(raw_row, mapping, "cost_currency")),
                inventory=_optional_int_text(_cell_from_mapping(raw_row, mapping, "inventory")),
                target_price=_optional_money_text(_cell_from_mapping(raw_row, mapping, "target_price")),
                notes=_cell_from_mapping(raw_row, mapping, "notes"),
            )
        )
    return parsed


def smart_rows_to_inventory_items(rows: list[SmartInventoryRow]) -> list[InventoryItem]:
    return [
        InventoryItem(
            product_name=row.product_name,
            cost=row.cost_price or 0.0,
            stock=row.inventory or 0,
            target_selling_price=row.target_price or 0.0,
            sku=row.sku,
            sku_count=row.inventory or 0,
            color=row.color,
            notes=row.notes,
            source="excel",
            inventory_unknown=row.inventory is None,
            cost_currency=row.cost_currency,
        )
        for row in rows
        if row.product_name.strip()
    ]


def smart_rows_to_inventory_text(rows: list[SmartInventoryRow]) -> str:
    lines = ["product_name,cost_price,inventory,target_selling_price,cost_currency"]
    for row in rows:
        name = _csv_safe(row.product_name)
        cost = "" if row.cost_price is None else row.cost_price
        inventory = "" if row.inventory is None else row.inventory
        target = "" if row.target_price is None else row.target_price
        lines.append(f"{name},{cost},{inventory},{target},{row.cost_currency}")
    return "\n".join(lines)


def _load_table_rows(file_bytes: bytes, filename: str) -> list[tuple[Any, ...]]:
    if filename.lower().endswith(".csv"):
        text = file_bytes.decode("utf-8-sig")
        return [tuple(row) for row in csv.reader(io.StringIO(text))]
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("仅支持 .xlsx 和 .csv 文件。")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Excel 上传需要安装 openpyxl：请运行 pip install openpyxl。") from exc

    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    return list(sheet.iter_rows(values_only=True))


def parse_taobao_inventory_json(raw_json: str) -> list[InventoryItem]:
    try:
        payload = json.loads(_normalize_taobao_payload(raw_json))
    except json.JSONDecodeError as exc:
        raise ValueError(f"淘宝 JSON 解析失败：第 {exc.lineno} 行第 {exc.colno} 列不是有效 JSON。") from exc

    products = parse_taobao_products(payload)
    return [
        InventoryItem(
            product_name=product.source_product_title,
            cost=product.price,
            stock=product.sku_number,
            target_selling_price=product.advise_sale_price_high or product.advise_sale_price_low or 0.0,
            sku=product.source_product_id,
            sku_count=product.sku_number,
            category=product.cat_name,
            status=product.target_product_status,
            source="taobao_json",
            cost_currency="CNY",
        )
        for product in products
    ]


def _normalize_taobao_payload(raw_text: str) -> str:
    text = raw_text.strip()
    if text.startswith("{"):
        return text
    match = re.match(r"^[\w$]+\(([\s\S]*)\)\s*;?$", text)
    if match:
        return match.group(1)
    raise ValueError("不支持的淘宝 payload 格式。请粘贴 JSON 或 mtopjsonp(... ) JSONP 响应。")


def _dedupe_inventory_items(items: list[InventoryItem]) -> list[InventoryItem]:
    deduped: dict[str, InventoryItem] = {}
    for item in items:
        key = item.sku.strip() or item.product_name.strip().lower()
        deduped[key] = item
    return list(deduped.values())


FIELD_KEYWORDS = {
    "product_name": ["商品名", "品名", "product", "title", "name", "款名", "名称"],
    "sku": ["sku", "货号", "编码", "款号", "商品编码", "货品编号"],
    "color": ["颜色", "color", "colour", "配色"],
    "cost_currency": ["成本币种", "币种", "currency", "cost_currency", "cost currency"],
    "cost_price": ["成本", "进价", "供货价", "cost", "采购价", "专属价"],
    "inventory": ["库存", "数量", "stock", "inventory", "可售", "件数"],
    "target_price": ["售价", "目标售价", "建议售价", "price", "直播价", "销售价", "定价"],
    "notes": ["备注", "卖点", "活动", "tag", "notes", "说明", "库存备注", "状态"],
}


def _detect_header_row(rows: list[tuple[Any, ...]]) -> int | None:
    best_index = None
    best_score = 0
    best_fields: set[str] = set()
    for index, row in enumerate(rows[:12]):
        headers = [_cell_text(value) for value in row]
        fields = {field for header in headers if (field := _field_from_header(header))}
        score = len(fields)
        if score > best_score:
            best_index = index
            best_score = score
            best_fields = fields
    if "product_name" in best_fields:
        return best_index
    if best_score >= 2 and "sku" in best_fields:
        return best_index
    return None


def _first_non_empty_row(rows: list[tuple[Any, ...]]) -> int | None:
    for index, row in enumerate(rows[:12]):
        if not _is_empty_row(row):
            return index
    return None


def _map_headers(headers: list[str], sample_rows: list[tuple[Any, ...]]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    used_columns: set[int] = set()
    for column_index, header in enumerate(headers):
        field = _field_from_header(header)
        if field and field not in mapping:
            mapping[field] = column_index
            used_columns.add(column_index)

    if "product_name" not in mapping:
        candidate = _guess_product_column(sample_rows, used_columns)
        if candidate is not None:
            mapping["product_name"] = candidate
            used_columns.add(candidate)
    for field in ("cost_price", "target_price"):
        if field not in mapping:
            candidate = _guess_numeric_column(sample_rows, used_columns, integer_only=(field == "inventory"))
            if candidate is not None:
                mapping[field] = candidate
                used_columns.add(candidate)
    return mapping


def _field_from_header(header: str) -> str | None:
    normalized = _normalize_text(header)
    for field, keywords in FIELD_KEYWORDS.items():
        if any(_normalize_text(keyword) in normalized for keyword in keywords):
            return field
    return None


def _normalize_currency(value: str) -> str:
    normalized = value.strip().upper()
    return normalized if normalized in {"CNY", "CAD", "USD"} else "CAD"


def _infer_cost_currency(cost_text: str, currency_text: str) -> str:
    explicit = currency_text.strip().upper()
    if explicit in {"CNY", "CAD", "USD"}:
        return explicit
    normalized = _normalize_text(cost_text)
    if re.search(r"[*x×]\s*[45]\.\d", normalized) or "人民币" in normalized or "cny" in normalized:
        return "CNY"
    return "CAD"


def _header_score(header: str) -> int:
    return 1 if _field_from_header(header) else 0


def _guess_product_column(sample_rows: list[tuple[Any, ...]], used_columns: set[int]) -> int | None:
    max_cols = max((len(row) for row in sample_rows), default=0)
    best: tuple[int, int] | None = None
    for column in range(max_cols):
        if column in used_columns:
            continue
        values = [_cell_text(row[column]) for row in sample_rows if column < len(row)]
        text_values = [value for value in values if value and not _looks_numeric(value)]
        score = sum(1 for value in text_values if len(value) >= 3)
        if best is None or score > best[0]:
            best = (score, column)
    return best[1] if best and best[0] > 0 else None


def _guess_numeric_column(
    sample_rows: list[tuple[Any, ...]],
    used_columns: set[int],
    integer_only: bool = False,
) -> int | None:
    max_cols = max((len(row) for row in sample_rows), default=0)
    best: tuple[int, int] | None = None
    for column in range(max_cols):
        if column in used_columns:
            continue
        values = [_cell_text(row[column]) for row in sample_rows if column < len(row)]
        numeric_count = 0
        for value in values:
            parsed = _optional_float_text(value)
            if parsed is None:
                continue
            if integer_only and parsed != int(parsed):
                continue
            numeric_count += 1
        if best is None or numeric_count > best[0]:
            best = (numeric_count, column)
    return best[1] if best and best[0] > 0 else None


def _parse_headerless_inventory_rows(rows: list[tuple[Any, ...]]) -> list[SmartInventoryRow]:
    data_rows = [row for row in rows if not _is_empty_row(row)]
    if not data_rows:
        return []

    mapping = _guess_headerless_columns(data_rows)
    product_column = mapping.get("product_name")
    if product_column is None:
        raise ValueError("没有识别到商品名列。支持表头：商品名、品名、product、title、name。")

    excluded_notes_columns = {column for column in mapping.values() if column is not None}
    parsed: list[SmartInventoryRow] = []
    for row in data_rows:
        product_name = _cell(row, product_column)
        if not product_name or _looks_like_date(product_name) or _looks_like_sku(product_name):
            continue
        notes = _joined_notes(row, excluded_notes_columns)
        cost_text = _cell(row, mapping.get("cost_price"))
        parsed.append(
            SmartInventoryRow(
                product_name=product_name,
                sku=_cell(row, mapping.get("sku")),
                color=_cell(row, mapping.get("color")),
                cost_price=_optional_money_text(cost_text),
                cost_currency=_infer_cost_currency(cost_text, ""),
                inventory=None,
                target_price=_optional_money_text(_cell(row, mapping.get("target_price"))),
                notes=notes,
            )
        )
    return parsed


def _guess_headerless_columns(rows: list[tuple[Any, ...]]) -> dict[str, int | None]:
    max_cols = max((len(row) for row in rows), default=0)
    sku_column = _guess_sku_column(rows, max_cols)
    product_column = _guess_headerless_product_column(rows, max_cols, {sku_column} if sku_column is not None else set())

    numeric_columns = [
        column
        for column in range(max_cols)
        if column not in {sku_column, product_column}
        and column > (product_column or -1)
        and _headerless_numeric_score(rows, column) > 0
    ]
    cost_column = numeric_columns[0] if numeric_columns else None
    target_column = numeric_columns[1] if len(numeric_columns) > 1 else None
    color_column = _guess_color_column(rows, product_column, sku_column, cost_column)

    return {
        "product_name": product_column,
        "sku": sku_column,
        "color": color_column,
        "cost_price": cost_column,
        "target_price": target_column,
    }


def _guess_sku_column(rows: list[tuple[Any, ...]], max_cols: int) -> int | None:
    best: tuple[int, int] | None = None
    for column in range(max_cols):
        score = sum(1 for row in rows if column < len(row) and _looks_like_sku(_cell_text(row[column])))
        if best is None or score > best[0]:
            best = (score, column)
    return best[1] if best and best[0] > 0 else None


def _guess_headerless_product_column(rows: list[tuple[Any, ...]], max_cols: int, used_columns: set[int]) -> int | None:
    best: tuple[int, int] | None = None
    for column in range(max_cols):
        if column in used_columns:
            continue
        score = 0
        for row in rows:
            value = _cell_text(row[column]) if column < len(row) else ""
            if not value or _looks_like_date(value) or _looks_like_sku(value) or _looks_numeric(value):
                continue
            score += _product_name_score(value)
        if best is None or score > best[0]:
            best = (score, column)
    return best[1] if best and best[0] > 0 else None


def _guess_color_column(
    rows: list[tuple[Any, ...]],
    product_column: int | None,
    sku_column: int | None,
    cost_column: int | None,
) -> int | None:
    if sku_column is not None and cost_column is not None and sku_column < cost_column:
        candidates = list(range(sku_column + 1, cost_column))
    elif product_column is not None and cost_column is not None and product_column < cost_column:
        candidates = list(range(product_column + 1, cost_column))
    else:
        candidates = []
    best: tuple[int, int] | None = None
    for column in candidates:
        score = 0
        for row in rows:
            value = _cell_text(row[column]) if column < len(row) else ""
            if value and not _looks_like_sku(value) and not _looks_numeric(value):
                score += 1
        if best is None or score > best[0]:
            best = (score, column)
    return best[1] if best and best[0] > 0 else None


def _headerless_numeric_score(rows: list[tuple[Any, ...]], column: int) -> int:
    score = 0
    for row in rows:
        value = _cell_text(row[column]) if column < len(row) else ""
        if not value or _looks_like_sku(value) or _looks_like_date(value):
            continue
        if _optional_money_text(value) is not None:
            score += 1
    return score


def _joined_notes(row: tuple[Any, ...], excluded_columns: set[int]) -> str:
    notes: list[str] = []
    for index, value in enumerate(row):
        text = _cell_text(value)
        if not text or index in excluded_columns:
            continue
        if _looks_like_sku(text) or re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text.replace(",", "")):
            continue
        notes.append(text)
    return " | ".join(notes)


def _product_name_score(value: str) -> int:
    normalized = _normalize_text(value)
    product_terms = [
        "arcteryx",
        "始祖鸟",
        "jacket",
        "hoody",
        "hoodie",
        "shirt",
        "pant",
        "pants",
        "vest",
        "coat",
        "fleece",
        "pullover",
        "crew",
        "neck",
        "shoe",
        "gtx",
        "atom",
        "beta",
        "cerium",
        "gamma",
        "kragg",
        "rho",
        "squamish",
        "emblem",
        "psiphon",
        "clarkia",
        "男",
        "女",
        "裤",
        "衣",
        "鞋",
    ]
    score = 1 if 3 <= len(value) <= 80 else 0
    score += sum(3 for term in product_terms if term in normalized)
    if any(term in normalized for term in ["库存", "更新", "折扣", "已做", "官网", "店内", "奥莱"]):
        score -= 2
    return max(score, 0)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_from_mapping(row: tuple[Any, ...], mapping: dict[str, int], field: str) -> str:
    index = mapping.get(field)
    if index is None or index >= len(row):
        return ""
    return _cell_text(row[index])


def _normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", value.strip().lower())


def _is_empty_row(row: tuple[Any, ...]) -> bool:
    return not any(_cell_text(value) for value in row)


def _looks_numeric(value: str) -> bool:
    return _optional_float_text(value) is not None


def _optional_float_text(value: str) -> float | None:
    cleaned = (
        str(value or "")
        .replace(",", "")
        .replace("$", "")
        .replace("¥", "")
        .replace("￥", "")
        .replace("（专属价）", "")
        .replace("专属价", "")
        .strip()
    )
    match = re.search(r"-?\d+(?:\.\d+)?", cleaned)
    if not match:
        return None
    try:
        return round(float(match.group(0)), 2)
    except ValueError:
        return None


def _optional_money_text(value: str) -> float | None:
    text = str(value or "").strip()
    if not text:
        return None
    cleaned = (
        text.replace(",", "")
        .replace("$", "")
        .replace("¥", "")
        .replace("￥", "")
        .replace("（专属价）", "")
        .replace("专属价", "")
        .strip()
    )
    expression_match = re.match(r"^\s*(-?\d+(?:\.\d+)?(?:\s*[+\-*/]\s*-?\d+(?:\.\d+)?)+)", cleaned)
    if expression_match:
        return _safe_eval_money_expression(expression_match.group(1))
    if re.search(r"(?:进价|成本|供货价)\s*[*x×]\s*\d", cleaned) and not re.match(r"^\s*-?\d", cleaned):
        return None
    return _optional_float_text(cleaned)


def _safe_eval_money_expression(expression: str) -> float | None:
    if not re.fullmatch(r"[\d.\s+\-*/]+", expression):
        return None
    try:
        value = eval(expression, {"__builtins__": {}}, {})
    except Exception:
        return None
    if not isinstance(value, (int, float)):
        return None
    return round(float(value), 2)


def _optional_int_text(value: str) -> int | None:
    if _looks_like_date(value):
        return None
    parsed = _optional_float_text(value)
    return int(parsed) if parsed is not None else None


def _looks_like_date(value: str) -> bool:
    text = str(value or "").strip()
    return bool(
        re.search(r"\b(?:19|20)\d{2}[-/.年]\d{1,2}(?:[-/.月]\d{1,2})?", text)
        or re.search(r"\b\d{1,2}[-/.]\d{1,2}[-/.](?:19|20)\d{2}\b", text)
    )


def _looks_like_sku(value: str) -> bool:
    return bool(re.search(r"\bX\d{6,}\b", str(value or "").strip(), re.IGNORECASE))


def _csv_safe(value: str) -> str:
    if any(char in value for char in [",", '"', "\n"]):
        return '"' + value.replace('"', '""') + '"'
    return value


def _cell(row: tuple[Any, ...], index: int | None) -> str:
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _float_cell(row: tuple[Any, ...], index: int | None, row_number: int, field_name: str) -> float:
    value = _optional_float_cell(row, index)
    if value is None:
        raise ValueError(f"Excel 第 {row_number} 行：{field_name} 必须是数字。")
    return value


def _optional_float_cell(row: tuple[Any, ...], index: int | None) -> float | None:
    text = _cell(row, index).replace(",", "").replace("$", "").replace("¥", "").replace("￥", "")
    if not text:
        return None
    try:
        return round(float(text), 2)
    except ValueError:
        return None


def _int_cell(row: tuple[Any, ...], index: int | None, row_number: int, field_name: str) -> int:
    value = _optional_float_cell(row, index)
    if value is None:
        raise ValueError(f"Excel 第 {row_number} 行：{field_name} 必须是整数。")
    return int(value)
